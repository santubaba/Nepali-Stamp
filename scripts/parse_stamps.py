import sys
import os
import json
import re
from docx import Document

def slugify(text):
    text = text.lower().strip()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[\s_]+', '-', text)
    text = re.sub(r'-+', '-', text)
    return text

SECTION_GROUPS = {
    "overview": [
        "the-1962-aaya-tekat-series",
        "beyond-the-postage-stamp",
        "nepals-1922-court-fee",
        "messages-across-the-clouds",
        "introduction",
    ],
    "genesisAndTimeline": [
        "genesis-production-and-administrative-timeline",
        "a-new-era-of-financial-remittance",
        "the-birth-of-airborne-correspondence",
    ],
    "legalFramework": [
        "the-legal-framework",
        "statutory-document-categories",
        "statutory-exemption",
        "legal-obligation-to-affix",
    ],
    "dutyRates": [
        "duty-rates-and-economic-applications",
        "distribution-and-vendor-commissions",
        "bank-vouchers-as-legal-equivalents",
    ],
    "anatomyAndIconography": [
        "anatomy-and-iconography",
        "physical-anatomy",
        "complete-breakdown-of-the-front",
        "reverse-side",
        "sender-information",
        "1-the-left-slip",
        "2-the-main-nepal-postal-order",
        "header-core-warning",
        "top-instructions",
        "denomination-panels",
        "redemption-receipt-section",
        "official-cancellation-rings",
    ],
    "localProduction": [
        "local-production",
        "proliferation-of-varieties",
    ],
    "judicialSecurity": [
        "judicial-security",
    ],
    "collectingSignificance": [
        "collecting-significance",
        "withdrawal-and-modern-collectibility",
        "philatelic-and-historical-legacy",
    ],
    "denominationOverview": [
        "denominations-color-tiers-and-paper-stocks",
        "denominations-and-color-schemes",
        "denomination-checklist-and-visual-iconography",
    ],
}

def get_image_extension(content_type):
    if content_type == "image/jpeg":
        return "jpg"

    if content_type == "image/png":
        return "png"

    if content_type == "image/gif":
        return "gif"

    if content_type == "image/webp":
        return "webp"

    return content_type.split("/")[-1]


def extract_image_from_cell(cell, images_dir, filename_base):
    """
    Extract the first image physically contained inside this table cell.

    Returns the filename if an image was found, otherwise None.
    """

    for paragraph in cell.paragraphs:
        for run in paragraph.runs:

            # Look for drawings/pictures in this run.
            drawing_elements = run._element.xpath(
                ".//*[local-name()='blip']"
            )

            for blip in drawing_elements:

                # Get the relationship ID used by this image.
                embed_id = blip.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )

                if not embed_id:
                    continue

                relationship = cell.part.rels.get(embed_id)

                if relationship is None:
                    continue

                image_part = relationship.target_part
                image_data = image_part.blob

                content_type = image_part.content_type
                extension = get_image_extension(content_type)

                filename = f"{filename_base}.{extension}"
                filepath = os.path.join(images_dir, filename)

                with open(filepath, "wb") as f:
                    f.write(image_data)

                return filename

    return None


def extract_sections(doc):
    sections = {}
    current_heading = "overview"
    current_paragraphs = []

    def flush():
        if current_paragraphs:
            text = " ".join(current_paragraphs).strip()

            if text:
                if current_heading in sections:
                    sections[current_heading] += " " + text
                else:
                    sections[current_heading] = text

    for para in doc.paragraphs:
        text = para.text.strip()

        if not text:
            continue

        is_heading = (
            para.style.name.startswith("Heading")
            or (
                len(text) < 120
                and all(run.bold for run in para.runs if run.text.strip())
                and para.runs
            )
        )

        if is_heading:
            flush()
            current_paragraphs = []
            current_heading = slugify(text)
        else:
            current_paragraphs.append(text)

    flush()

    return sections


def group_sections(raw_sections):
    grouped = {}
    used_keys = set()

    for group_key, prefixes in SECTION_GROUPS.items():
        combined = []

        for raw_key in raw_sections:
            for prefix in prefixes:
                if raw_key.startswith(prefix) and raw_key not in used_keys:
                    combined.append(raw_sections[raw_key])
                    used_keys.add(raw_key)
                    break

        if combined:
            grouped[group_key] = " ".join(combined)

    for raw_key, value in raw_sections.items():
        if raw_key not in used_keys:
            grouped[raw_key] = value

    return grouped


def extract_production_and_issuance(sections):
    production = {}
    issuance = {}

    full_text = " ".join(sections.values())

    printer_match = re.search(
        r'Printer[:\s]+(.+?)(?:\s{2,}|\*\*\n\*\*|First Printed|Printing)',
        full_text
    )

    if printer_match:
        production["printer"] = printer_match.group(1).strip().rstrip(".")

    method_match = re.search(
        r'Printing Technique[:\s]+(.+?)(?:\s{2,}|\*\*\n\*\*|Paper|First)',
        full_text
    )

    if not method_match:
        method_match = re.search(
            r'Printing Method[:\s]+(.+?)(?:\s{2,}|\*\*\n\*\*|Paper|First)',
            full_text
        )

    if method_match:
        production["method"] = method_match.group(1).strip().rstrip(".")

    paper_match = re.search(
        r'Paper Stock[:\s]+(.+?)(?:\s{2,}|\*\*\n\*\*|First|Issued)',
        full_text
    )

    if not paper_match:
        paper_match = re.search(
            r'Paper[:\s]+(.+?)(?:\s{2,}|\*\*\n\*\*|First|Issued)',
            full_text
        )

    if paper_match:
        production["paper"] = paper_match.group(1).strip().rstrip(".")

    first_printed_match = re.search(
        r'First Printed[:\s]+(.+?)(?:\s{2,}|\*\*\n\*\*|Issued)',
        full_text
    )

    if first_printed_match:
        issuance["firstPrinted"] = (
            first_printed_match.group(1).strip().rstrip(".")
        )

    circulation_match = re.search(
        r'Issued into Circulation[:\s]+(.+?)(?:—|\s{2,}|\*\*\n\*\*|Demonetized)',
        full_text
    )

    if circulation_match:
        issuance["issuedIntoCirculation"] = (
            circulation_match.group(1).strip().rstrip(".")
        )

    demonetized_match = re.search(
        r'Demonetized[^:]*:\s*(.+?\d{4}\s*(?:AD)?[).])',
        full_text
    )

    if demonetized_match:
        issuance["demonetized"] = (
            demonetized_match.group(1).strip().rstrip(".")
        )

    return (
        production if production else None,
        issuance if issuance else None,
    )


def extract_eyebrow(doc):
    for para in doc.paragraphs:
        text = para.text.strip()

        if not text:
            continue

        match = re.search(
            r'(\d{4}[^:,\n]{0,60}(?:Series|Issue|Orders|Order|Stamps|Aerogrammes))',
            text,
            re.IGNORECASE,
        )

        if match:
            return match.group(1).strip()

        match = re.search(
            r"Nepal[\u2019']s\s+([^(]+)\s*\(?\d{4}",
            text,
            re.IGNORECASE,
        )

        if match:
            return match.group(1).strip()

        return text[:80]

    return ""


def find_denomination_table(doc):
    best_table = None
    best_row_count = 0

    for table in doc.tables:

        if len(table.rows) < 2:
            continue

        headers = [
            cell.text.strip().lower()
            for cell in table.rows[0].cells
        ]

        priority_keywords = [
            "denomination",
            "inscription",
            "denomination (face value)",
        ]

        fallback_keywords = [
            "value",
            "face",
            "denom",
            "year",
            "value group",
        ]

        is_priority = any(
            any(kw in header for kw in priority_keywords)
            for header in headers
        )

        is_fallback = any(
            any(kw in header for kw in fallback_keywords)
            for header in headers
        )

        if is_priority:
            if len(table.rows) > best_row_count:
                best_table = table
                best_row_count = len(table.rows)

        elif is_fallback and best_table is None:
            best_table = table
            best_row_count = len(table.rows)

    return best_table


def parse_table_row(
    headers,
    row,
    category_slug,
    images_dir,
    row_index,
    eyebrow,
    sections,
    production,
    issuance,
):
    """
    Parse one denomination row.

    IMPORTANT:
    The image is extracted directly from this row,
    rather than relying on global DOCX image order.
    """

    row_cells = row.cells
    data = {}

    for i, cell in enumerate(row_cells):
        if i < len(headers):
            data[headers[i]] = cell.text.strip()

    title = (
        data.get("denomination")
        or data.get("face value")
        or data.get("denomination (face value)")
        or data.get("value")
        or data.get("year")
        or data.get("issue / subject")
        or f"Record {row_index + 1}"
    )

    if not title or title.lower() in ["photo", ""]:
        return None

    slug = f"{category_slug}-{slugify(title)}"

    # --------------------------------------------------
    # Extract image from THIS TABLE ROW
    # --------------------------------------------------

    image = None

    for cell_index, cell in enumerate(row_cells):

        # Search every cell because we don't yet assume
        # which column contains the image.
        filename_base = slug

        extracted = extract_image_from_cell(
            cell,
            images_dir,
            filename_base,
        )

        if extracted:
            image = extracted
            print(
                f"    Image found for {title}: "
                f"{extracted} "
                f"(cell {cell_index})"
            )
            break

    if not image:
        print(f"    WARNING: No image found for {title}")

    # --------------------------------------------------
    # Tags
    # --------------------------------------------------

    tags = {}

    for key in [
        "color",
        "colour",
        "color & paper",
        "color tiers",
        "color specifications",
        "color & paper specifications",
    ]:
        if key in data and data[key]:
            tags["color"] = data[key]
            break

    # --------------------------------------------------
    # Key attributes
    # --------------------------------------------------

    key_attributes = {
        "denomination": title
    }

    for key in [
        "primary motif & iconography",
        "motif",
        "primary usage",
        "primary role",
        "inscription in plug",
    ]:
        if key in data and data[key]:
            key_attributes["motif"] = data[key]
            break

    for key in ["distinguishing features"]:
        if key in data and data[key]:
            key_attributes["distinguishingFeatures"] = data[key]

    for key in ["face values", "face value"]:
        if key in data and data[key]:
            key_attributes["faceValues"] = data[key]
            break

    # --------------------------------------------------
    # Physical properties
    # --------------------------------------------------

    physical = {}

    for key in ["stamp dimensions", "dimensions"]:
        if key in data and data[key]:
            physical["dimensions"] = data[key]
            break

    if "perforation" in data and data["perforation"]:
        physical["perforation"] = data["perforation"]

    # --------------------------------------------------
    # Final stamp record
    # --------------------------------------------------

    return {
        "title": title,
        "slug": slug,
        "eyebrow": eyebrow,
        "image": image or "placeholder.jpg",
        "featured": False,
        "tags": tags if tags else None,
        "keyAttributes": key_attributes if key_attributes else None,
        "physicalProperties": physical if physical else None,
        "printingProduction": production,
        "issuance": issuance,
        "historicalContext": sections if sections else None,
    }


def parse_word_file(filepath, category_slug):

    doc = Document(filepath)

    base_dir = os.path.dirname(filepath)

    images_dir = os.path.join(
        base_dir,
        "images",
    )

    output_dir = os.path.join(
        base_dir,
        "output",
    )

    os.makedirs(images_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)

    print("Extracting sections...")

    raw_sections = extract_sections(doc)
    sections = group_sections(raw_sections)

    print(
        f"  Found sections: {list(sections.keys())}"
    )

    eyebrow = extract_eyebrow(doc)

    print(
        f"  Eyebrow: {eyebrow}"
    )

    production, issuance = extract_production_and_issuance(
        sections
    )

    print(
        f"  Production: {production}"
    )

    print(
        f"  Issuance: {issuance}"
    )

    print("Finding denomination table...")

    table = find_denomination_table(doc)

    if not table:
        print(
            "  ERROR: No denomination table found in document"
        )
        sys.exit(1)

    headers = [
        cell.text.strip().lower()
        for cell in table.rows[0].cells
    ]

    print(
        f"  Table headers: {headers}"
    )

    stamps = []

    # --------------------------------------------------
    # Parse each table row
    # --------------------------------------------------

    for row_index, row in enumerate(table.rows[1:]):

        cells = [
            cell.text.strip()
            for cell in row.cells
        ]

        if not any(cells):
            continue

        stamp = parse_table_row(
            headers=headers,
            row=row,
            category_slug=category_slug,
            images_dir=images_dir,
            row_index=row_index,
            eyebrow=eyebrow,
            sections=sections,
            production=production,
            issuance=issuance,
        )

        if stamp is None:
            continue

        stamps.append(stamp)

        print(
            f"  Parsed: {stamp['title']} "
            f"→ {stamp['image']}"
        )

    # --------------------------------------------------
    # Write JSON
    # --------------------------------------------------

    output_path = os.path.join(
        output_dir,
        f"{category_slug}.json",
    )

    with open(
        output_path,
        "w",
        encoding="utf-8",
    ) as f:

        json.dump(
            {
                "category": category_slug,
                "stamps": stamps,
            },
            f,
            ensure_ascii=False,
            indent=2,
        )

    print(
        f"\nDone. {len(stamps)} stamps written to:"
    )

    print(output_path)


if __name__ == "__main__":

    if len(sys.argv) != 3:

        print(
            "Usage: python parse_stamps.py "
            "<path-to-docx> <category-slug>"
        )

        print(
            "Example: python parse_stamps.py "
            "scripts/aerogrammes.docx aerogrammes"
        )

        sys.exit(1)

    filepath = sys.argv[1]
    category_slug = sys.argv[2]

    if not os.path.exists(filepath):

        print(
            f"ERROR: File not found: {filepath}"
        )

        sys.exit(1)

    parse_word_file(
        filepath,
        category_slug,
    )